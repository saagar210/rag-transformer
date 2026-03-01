import io
import pytest
from fastapi import HTTPException

from app import _TextExtractor, _extract_docx, extract_text


class TestHtmlExtractor:
    def test_extracts_body_text(self):
        p = _TextExtractor()
        p.feed("<p>Hello world</p>")
        assert "Hello world" in p.get_text()

    def test_strips_script_tags(self):
        p = _TextExtractor()
        p.feed("<script>var x=1;</script>Hello")
        assert "var x" not in p.get_text()
        assert "Hello" in p.get_text()

    def test_strips_style_tags(self):
        p = _TextExtractor()
        p.feed("<style>.a{color:red}</style>Text")
        assert "color" not in p.get_text()
        assert "Text" in p.get_text()

    def test_nested_script_depth(self):
        # Malformed HTML with nested script — depth counter prevents bleed
        p = _TextExtractor()
        p.feed("<script><script>inner</script>between</script>after")
        text = p.get_text()
        assert "inner" not in text
        assert "after" in text

    def test_sample_html_fixture(self):
        with open("tests/fixtures/sample.html", "rb") as f:
            content = f.read()
        p = _TextExtractor()
        p.feed(content.decode())
        text = p.get_text()
        assert "Main heading" in text
        assert "Body paragraph" in text
        assert "color: red" not in text
        assert "var x" not in text


class TestDocxExtractor:
    def test_extracts_paragraphs(self, sample_docx_bytes):
        text = _extract_docx(sample_docx_bytes)
        assert "First paragraph" in text
        assert "Second paragraph" in text

    def test_extracts_table_cells(self):
        from docx import Document
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        table.cell(1, 0).text = "C"
        table.cell(1, 1).text = "D"
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        text = _extract_docx(buf.getvalue())
        assert "A" in text and "B" in text


class TestExtractText:
    def test_txt(self):
        assert extract_text("doc.txt", b"hello") == "hello"

    def test_md(self):
        assert extract_text("doc.md", b"# Title\nBody") == "# Title\nBody"

    def test_rst(self):
        assert "content" in extract_text("doc.rst", b"content")

    def test_html(self):
        text = extract_text("page.html", b"<p>Hi</p>")
        assert "Hi" in text

    def test_docx(self, sample_docx_bytes):
        text = extract_text("doc.docx", sample_docx_bytes)
        assert "First paragraph" in text

    def test_pdf(self, sample_pdf_bytes):
        text = extract_text("doc.pdf", sample_pdf_bytes)
        assert "Hello from PDF" in text

    def test_unsupported_raises_422(self):
        with pytest.raises(HTTPException) as exc:
            extract_text("file.exe", b"data")
        assert exc.value.status_code == 422
        assert ".exe" in exc.value.detail
